"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu.

Tên cột sheet "Experiments" (giữ nguyên thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính — KHÔNG ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl

# 4 cột công thức của mẫu: bỏ qua khi ghi
FORMULA_COLS = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}

OPT_LABEL = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}
LOSS_LABEL = {"ce": "CE", "mse": "MSE"}


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi cfg + history + summary (KHÔNG ghi best_state) ra <results_dir>/<exp_id>.json."""
    d = Path(results_dir)
    d.mkdir(parents=True, exist_ok=True)
    path = d / f"{result['cfg']['exp_id']}.json"
    payload = dict(cfg=result["cfg"], history=result["history"], summary=result["summary"])
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, ensure_ascii=False, default=float)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    files = sorted(Path(results_dir).glob("*.json"))
    return [json.load(open(p, encoding="utf-8")) for p in files]


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng (khoá trùng tên cột ở đầu file).

    eval_scores: chỉ truyền cho baseline và cấu hình cuối cùng, dạng {"accuracy": .., "macro_f1": ..}
                 lấy từ eval_result.json (do scripts/evaluate.py tạo), KHÔNG tự tính lại.
    """
    cfg, s = result["cfg"], result["summary"]
    row = dict(
        exp_id=cfg["exp_id"], group=cfg["group"], description=cfg["description"],
        loss=LOSS_LABEL.get(cfg["loss"], cfg["loss"]),
        optimizer=OPT_LABEL.get(cfg["optimizer"], cfg["optimizer"]),
        lr=cfg["lr"], weight_decay=cfg["weight_decay"], batch=cfg["batch"], epochs=cfg["epochs"],
        hidden="-".join(map(str, cfg["hidden"])), dropout=cfg["dropout"],
        clip_norm="none" if cfg["clip_norm"] is None else cfg["clip_norm"],
        precision=cfg["precision"], init=cfg["init"], seed=cfg["seed"],
        step0_loss=s["step0_loss"], best_val_loss=s["best_val_loss"], best_epoch=s["best_epoch"],
        final_train_loss=s["final_train_loss"], final_val_loss=s["final_val_loss"],
        val_acc=s["val_acc"], val_macro_f1=s["val_macro_f1"],
        time_per_epoch_s=s["time_per_epoch_s"], peak_mem_MB=(s["peak_mem_MB"] or None),
        diverged="Y" if s["diverged"] else "N",
        figure_file=f"figures/{cfg['exp_id']}.png", notes=notes,
    )
    if eval_scores:
        row["eval_acc"] = eval_scores.get("accuracy")
        row["eval_macro_f1"] = eval_scores.get("macro_f1")
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seeds_exp_ids: list[str] | None = None,
               summary_notes: dict | None = None) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu (từ dòng 2), lưu thành out_path.

    seeds_exp_ids : exp_id các lần chạy baseline khác seed -> ghi vào sheet "Seeds" cột A
                    (sheet này tự tính mean / σ / 2σ).
    summary_notes : {group: "nhận xét"} -> ghi vào cột "nhận xét ngắn" của sheet "Summary".
    Giữ nguyên công thức của mẫu (không dùng data_only=True).
    """
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]
    header = [c.value for c in ws[1]]
    col_of = {name: i + 1 for i, name in enumerate(header) if name}

    unknown = set(rows[0].keys()) - set(col_of) if rows else set()
    assert not unknown, f"có khoá không khớp cột của mẫu: {unknown}"

    for r, row in enumerate(rows, start=2):
        for key, val in row.items():
            if key in FORMULA_COLS:                # ô xám: công thức của mẫu, không ghi đè
                continue
            ws.cell(row=r, column=col_of[key], value=val)

    # ---- sheet Seeds: điền exp_id các lần chạy baseline (công thức tự tra Experiments) ----
    if seeds_exp_ids:
        wss = wb["Seeds"]
        for i, eid in enumerate(seeds_exp_ids):
            wss.cell(row=2 + i, column=1, value=eid)

    # ---- sheet Summary: nhận xét ngắn theo nhóm ----
    if summary_notes:
        wsm = wb["Summary"]
        head_s = [c.value for c in wsm[1]]
        col_note = [i + 1 for i, v in enumerate(head_s) if v and "nhận xét" in str(v)]
        if col_note:
            for r in range(2, wsm.max_row + 1):
                g = wsm.cell(row=r, column=1).value
                if g in summary_notes:
                    wsm.cell(row=r, column=col_note[0], value=summary_notes[g])

    wb.save(out_path)
    print(f"đã ghi {out_path}: {len(rows)} dòng thí nghiệm"
          + (f", {len(seeds_exp_ids)} seed baseline" if seeds_exp_ids else ""))
